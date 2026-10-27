"""生成分类标注表：从 private/ 的真实账单里随机抽 100 笔支出，让人填「正确分类」。

运行：uv run python scripts/make_label_sheet.py
输出：private/分类标注表.xlsx（在 private/ 里，不会上传到 GitHub）

注意：
  - 表里故意不显示程序的分类结果，免得影响标注的人
  - 随机种子固定，每次抽到的是同一批交易
  - 如果标注表已经填了内容，不会覆盖（加 --force 才会重新生成）
"""

import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from common import LABEL_SHEET, load_private_bills
from core.schema import CATEGORIES

SAMPLE_SIZE = 100
SEED = 2026
LABEL_COLUMN = "正确分类"
HEADERS = ["编号", "时间", "平台", "交易对方", "商品", "金额", "支付方式", "平台分类", LABEL_COLUMN, "备注（可不填）"]


def already_labelled() -> int:
    if not LABEL_SHEET.exists():
        return 0
    ws = load_workbook(LABEL_SHEET, read_only=True)["标注"]
    col = HEADERS.index(LABEL_COLUMN)
    return sum(1 for row in ws.iter_rows(min_row=2, values_only=True) if row[col])


def main():
    filled = already_labelled()
    if filled and "--force" not in sys.argv:
        sys.exit(f"标注表里已经填了 {filled} 笔，为了不覆盖你的标注，没有重新生成。确实要重新生成请加 --force")

    df = load_private_bills()
    spend = df[df["direction"] == "支出"]
    sample = spend.sample(n=min(SAMPLE_SIZE, len(spend)), random_state=SEED).sort_values("time")

    wb = Workbook()
    ws = wb.active
    ws.title = "标注"
    ws.append(HEADERS)
    for _, r in sample.iterrows():
        ws.append([r["id"], r["time"].strftime("%Y-%m-%d %H:%M"), r["platform"], r["counterparty"], r["item"],
                   round(float(r["amount"]), 2), r["pay_method"], r["raw_type"], None, None])

    # 表头样式、冻结首行、列宽
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFF2CC" if cell.value == LABEL_COLUMN else "EEEEEE")
    ws.freeze_panes = "B2"
    for col, width in zip("ABCDEFGHIJ", [6, 17, 7, 22, 30, 9, 18, 14, 12, 20]):
        ws.column_dimensions[col].width = width
    ws.column_dimensions["A"].hidden = True   # 编号只给程序用，藏起来

    # 「正确分类」这一列做成下拉框
    label_col = "ABCDEFGHIJ"[HEADERS.index(LABEL_COLUMN)]
    dv = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="分类不对", error="请从下拉框里选一个分类")
    ws.add_data_validation(dv)
    dv.add(f"{label_col}2:{label_col}{len(sample) + 1}")

    guide = wb.create_sheet("说明")
    for line in [
        "怎么填：",
        f"1. 在「标注」表的「{LABEL_COLUMN}」那一列（黄色表头），点格子右边的小箭头，选这笔钱实际属于哪个分类。",
        "2. 按你自己的真实情况判断，比如转账给同学是还饭钱，也请选「社交」。",
        "3. 实在判断不了的，选「其他」，并在备注里写一句为什么。",
        "4. 填完保存，告诉 Claude 就行。可以分几次填完。",
        "",
        "分类说明（按「买了什么」分，不按「在哪买」分）：",
        "餐饮：食堂、外卖、餐馆　　零食饮品：奶茶、咖啡、零食、便利店、水果",
        "购物：网购、超市、日用品、衣服、数码　　交通：公交地铁、打车、火车票",
        "娱乐：电影、游戏、网吧、视频和音乐会员　　生活缴费：话费、电费、水费",
        "学习：教材、打印、考试报名、网课　　医疗：药店、医院",
        "社交：转账、红包、AA　　其他：以上都不是",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 100
    for row in guide.iter_rows():
        row[0].alignment = Alignment(wrap_text=False)

    LABEL_SHEET.parent.mkdir(exist_ok=True)
    wb.save(LABEL_SHEET)
    print(f"已生成标注表：{LABEL_SHEET}（{len(sample)} 笔）")


if __name__ == "__main__":
    main()
